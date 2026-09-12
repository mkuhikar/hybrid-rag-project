import io
import pandas as pd
import docx
from pypdf import PdfReader
from pptx import Presentation
from typing import List, Dict, Any

def extract_text_from_pdf(file_obj) -> str:
    file_obj.seek(0)
    reader = PdfReader(file_obj)
    text = []
    for page in reader.pages:
        extracted = page.extract_text()
        if extracted:
            text.append(extracted)
    return "\n".join(text)

def extract_text_from_docx(file_obj) -> str:
    file_obj.seek(0)
    docx_bytes = io.BytesIO(file_obj.read())
    doc = docx.Document(docx_bytes)
    
    full_text = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_text:
                full_text.append(row_text)
                
    return "\n\n".join(full_text)

def extract_text_from_excel(file_obj) -> str:
    file_obj.seek(0)
    excel_file = pd.ExcelFile(file_obj)
    extracted_text = []
    
    for sheet_name in excel_file.sheet_names:
        df = pd.read_excel(file_obj, sheet_name=sheet_name)
        extracted_text.append(f"--- Sheet: {sheet_name} ---")
        extracted_text.append(df.to_csv(index=False, sep="\t"))
        
    return "\n\n".join(extracted_text)

def extract_text_from_pptx(file_obj) -> str:
    file_obj.seek(0)
    prs = Presentation(file_obj)
    slides_text = []
    
    for idx, slide in enumerate(prs.slides, start=1):
        slide_lines = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_lines.append(shape.text.strip())
        if slide_lines:
            slides_text.append(f"--- Slide {idx} ---\n" + "\n".join(slide_lines))
            
    return "\n\n".join(slides_text)

def parse_uploaded_files(uploaded_files_dict: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Converts uploaded file buffers into structured document dictionaries for RAG indexing."""
    parsed_docs = []
    
    for filename, file_obj in uploaded_files_dict.items():
        ext = filename.split(".")[-1].lower()
        extracted_text = ""
        
        try:
            if ext == "pdf":
                extracted_text = extract_text_from_pdf(file_obj)
            elif ext in ["docx", "doc"]:
                extracted_text = extract_text_from_docx(file_obj)
            elif ext == "xlsx":
                extracted_text = extract_text_from_excel(file_obj)
            elif ext == "pptx":
                extracted_text = extract_text_from_pptx(file_obj)
            
            if extracted_text.strip():
                parsed_docs.append({
                    "id": filename,
                    "text": extracted_text
                })
        except Exception as e:
            print(f"Error parsing file {filename}: {str(e)}")
            
    return parsed_docs